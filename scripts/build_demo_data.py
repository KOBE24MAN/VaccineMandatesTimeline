"""Build the immutable demo CSV, SQLite database, and release manifest."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import sqlite3
import sys
import tempfile
from collections import Counter
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.database import create_database_engine
from scripts.init_database import MANDATE_COLUMNS, initialize_database


SOURCE_COLUMNS = MANDATE_COLUMNS[:-2]
DATE_COLUMNS = ("effective_date", "enforcement_date", "removal_date")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def normalize_date(value: object, *, field: str, mandate_id: str) -> str:
    if value is None or clean_text(value) == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    raw = clean_text(value)
    for parser in (
        lambda candidate: datetime.fromisoformat(candidate).date(),
        lambda candidate: datetime.strptime(candidate, "%d/%m/%Y").date(),
        lambda candidate: datetime.strptime(candidate, "%Y/%m/%d").date(),
    ):
        try:
            return parser(raw).isoformat()
        except ValueError:
            continue
    raise ValueError(f"Invalid {field} for {mandate_id}: {raw}")


def normalize_boolean(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return clean_text(value).upper() in {"1", "TRUE", "YES", "Y"}


def normalize_integer(value: object) -> int | None:
    if value is None or clean_text(value) == "":
        return None
    return int(float(value))


def read_visibility_seed(path: Path) -> dict[str, int]:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        result: dict[str, int] = {}
        for row in reader:
            mandate_id = clean_text(row.get("id"))
            raw_level = clean_text(row.get("visibility_level"))
            if mandate_id and raw_level:
                level = int(float(raw_level))
                if 1 <= level <= 6:
                    result[mandate_id] = level
        return result


def derive_visibility(duration_days: int | None) -> int:
    if duration_days is None:
        return 6
    if duration_days >= 365:
        return 1
    if duration_days >= 270:
        return 2
    if duration_days >= 180:
        return 3
    if duration_days >= 90:
        return 4
    if duration_days >= 30:
        return 5
    return 6


def read_workbook(source_path: Path, visibility_seed: Path) -> list[dict[str, object]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    if headers != SOURCE_COLUMNS:
        raise ValueError(
            "Workbook columns do not match the required schema. "
            f"Expected {SOURCE_COLUMNS}, received {headers}"
        )

    visibility = read_visibility_seed(visibility_seed)
    records: list[dict[str, object]] = []
    seen_ids: set[str] = set()
    for row_number, values in enumerate(rows, start=2):
        source_row = dict(zip(headers, values, strict=True))
        mandate_id = clean_text(source_row["id"])
        if not mandate_id:
            raise ValueError(f"Missing mandate ID at workbook row {row_number}")
        if mandate_id in seen_ids:
            raise ValueError(f"Duplicate mandate ID in workbook: {mandate_id}")
        seen_ids.add(mandate_id)

        jurisdiction = clean_text(source_row["jurisdiction"])
        if not jurisdiction:
            raise ValueError(f"Missing jurisdiction for {mandate_id}")

        record: dict[str, object] = {
            column: clean_text(source_row[column]) for column in SOURCE_COLUMNS
        }
        record["id"] = mandate_id
        record["jurisdiction"] = jurisdiction
        for field in DATE_COLUMNS:
            record[field] = normalize_date(
                source_row[field], field=field, mandate_id=mandate_id
            )
        duration = normalize_integer(source_row["duration_days"])
        record["duration_days"] = "" if duration is None else duration
        record["date_uncertain"] = normalize_boolean(source_row["date_uncertain"])
        record["ongoing"] = not bool(record["removal_date"])
        record["visibility_level"] = visibility.get(
            mandate_id, derive_visibility(duration)
        )
        records.append(record)

    workbook.close()
    if not records:
        raise ValueError("Workbook contains no mandate records")
    return records


def write_csv(path: Path, records: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(
            destination,
            fieldnames=MANDATE_COLUMNS,
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(records)


def database_summary(path: Path) -> dict[str, object]:
    with closing(sqlite3.connect(path)) as connection:
        total = connection.execute("SELECT COUNT(*) FROM mandates").fetchone()[0]
        distinct_ids = connection.execute(
            "SELECT COUNT(DISTINCT id) FROM mandates"
        ).fetchone()[0]
        uncertain = connection.execute(
            "SELECT COUNT(*) FROM mandates WHERE date_uncertain = 1"
        ).fetchone()[0]
        ongoing = connection.execute(
            "SELECT COUNT(*) FROM mandates WHERE ongoing = 1"
        ).fetchone()[0]
        notable_events = connection.execute(
            "SELECT COUNT(*) FROM notable_events"
        ).fetchone()[0]
        categories = connection.execute(
            "SELECT COUNT(*) FROM mandate_categories"
        ).fetchone()[0]
        date_range = connection.execute(
            "SELECT MIN(COALESCE(effective_date, enforcement_date)), "
            "MAX(COALESCE(removal_date, effective_date, enforcement_date)) "
            "FROM mandates"
        ).fetchone()
        jurisdiction_rows = connection.execute(
            "SELECT jurisdiction, COUNT(*) FROM mandates "
            "GROUP BY jurisdiction ORDER BY jurisdiction"
        ).fetchall()
    return {
        "records": total,
        "distinct_ids": distinct_ids,
        "date_uncertain": uncertain,
        "ongoing": ongoing,
        "notable_events": notable_events,
        "category_links": categories,
        "date_range": {"earliest": date_range[0], "latest": date_range[1]},
        "by_jurisdiction": {row[0]: row[1] for row in jurisdiction_rows},
    }


def build_release(
    *,
    source_path: Path,
    visibility_seed: Path,
    events_path: Path,
    output_dir: Path,
    version: str,
) -> dict[str, object]:
    source_path = source_path.resolve()
    visibility_seed = visibility_seed.resolve()
    events_path = events_path.resolve()
    output_dir = output_dir.resolve()
    for required_path in (source_path, visibility_seed, events_path):
        if not required_path.is_file():
            raise FileNotFoundError(f"Required input does not exist: {required_path}")

    records = read_workbook(source_path, visibility_seed)
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mandeval-build-") as temp_name:
        temp_dir = Path(temp_name)
        csv_path = temp_dir / "mandates.csv"
        database_path = temp_dir / "mandates.db"
        write_csv(csv_path, records)

        engine = create_database_engine(f"sqlite:///{database_path.as_posix()}")
        initialized = initialize_database(engine, csv_path, events_path)
        engine.dispose()
        with closing(sqlite3.connect(database_path)) as connection:
            connection.execute("VACUUM")

        summary = database_summary(database_path)
        if summary["records"] != len(records):
            raise RuntimeError("Database record count does not match the workbook")
        if summary["distinct_ids"] != len(records):
            raise RuntimeError("Database contains duplicate mandate IDs")

        final_csv = output_dir / "mandates.csv"
        final_database = output_dir / "mandates.db"
        shutil.copyfile(csv_path, final_csv)
        shutil.copyfile(database_path, final_database)

        manifest: dict[str, object] = {
            "manifest_version": 1,
            "release_version": version,
            "source": {
                "file_name": source_path.name,
                "sha256": sha256_file(source_path),
                "records": len(records),
            },
            "inputs": {
                "visibility_seed_sha256": sha256_file(visibility_seed),
                "notable_events_sha256": sha256_file(events_path),
            },
            "outputs": {
                "mandates_csv": {
                    "file_name": final_csv.name,
                    "sha256": sha256_file(final_csv),
                },
                "mandates_database": {
                    "file_name": final_database.name,
                    "sha256": sha256_file(final_database),
                },
            },
            "validation": {
                **summary,
                "date_uncertain_ids": [
                    str(record["id"])
                    for record in records
                    if bool(record["date_uncertain"])
                ],
                "negative_duration_ids": [
                    str(record["id"])
                    for record in records
                    if isinstance(record["duration_days"], int)
                    and int(record["duration_days"]) < 0
                ],
                "missing_start_date_ids": [
                    str(record["id"])
                    for record in records
                    if not record["effective_date"] and not record["enforcement_date"]
                ],
                "visibility_levels": dict(
                    sorted(
                        Counter(
                            int(record["visibility_level"]) for record in records
                        ).items()
                    )
                ),
                "database_insert_counts": initialized,
            },
        }
        manifest_path = output_dir / "release-manifest.json"
        manifest_path.write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build validated demo data from the approved Excel workbook"
    )
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument(
        "--visibility-seed", type=Path, default=PROJECT_ROOT / "All_Mandates.csv"
    )
    parser.add_argument(
        "--events", type=Path, default=PROJECT_ROOT / "data" / "notable_events.csv"
    )
    parser.add_argument(
        "--output-dir", type=Path, default=PROJECT_ROOT / "data" / "release"
    )
    parser.add_argument("--version", default="demo-2026-08-26")
    args = parser.parse_args()

    manifest = build_release(
        source_path=args.source,
        visibility_seed=args.visibility_seed,
        events_path=args.events,
        output_dir=args.output_dir,
        version=args.version,
    )
    validation = manifest["validation"]
    print(
        "[OK] Built demo data: "
        f"{validation['records']} mandates, "
        f"{validation['date_uncertain']} uncertain dates, "
        f"{validation['notable_events']} notable events"
    )


if __name__ == "__main__":
    main()
